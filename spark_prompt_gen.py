#!/usr/bin/env python3
"""
Spark数据源提示词生成工具
基于spark_simple.py连接Spark集群，生成模板文件和prompt文件
"""

from prompt import tool, loader, gen
from server import conf, aws, llm
from server.db import spark as spark_db
import os
import json
import argparse
from openpyxl import Workbook

conf.load_env()

bucket_name = os.getenv("BUCKET_NAME")
example_file = os.getenv("EXAMPLE_FILE_NAME")
prompt_file = os.getenv("PROMPT_FILE_NAME")
rag_file = os.getenv("RAG_FILE_NAME")

data_files = f"{os.getcwd()}/prompt/data/promptdata"
save_to_path = f"{os.getcwd()}/prompt/prompt_conf"

prompt_file_name = prompt_file.split("/")[1]
example_file_name = example_file.split("/")[1]
rag_file_name = rag_file.split("/")[1]

prompt_path = save_to_path + "/" + prompt_file_name
example_path = save_to_path + "/" + example_file_name
rag_path = save_to_path + "/" + rag_file_name


def confirm_action(prompt_input):
    """提示用户确认操作"""
    response = input(prompt_input).strip().lower()
    return response in ['y', 'yes']


def get_spark_table_schema(table_name, conn):
    """获取Spark表结构（兼容gen.py接口）"""
    # 如果表名包含数据库前缀，分离出来
    if '.' in table_name:
        database, table = table_name.split('.', 1)
        schema_info = conn.get_table_schema(table, database)
    else:
        schema_info = conn.get_table_schema(table_name)
    
    schema = []
    for col in schema_info:
        schema.append({
            'Name': col['column_name'],
            'Type': col['data_type'],
            'Key': '',
            'Comment': col.get('comment', '')
        })
    return schema


def get_spark_sample_data(table_name, columns, conn, sample_count=3):
    """获取Spark表样本数据（兼容gen.py接口）"""
    column_names = [col['Name'] for col in columns]
    columns_str = ', '.join(column_names)
    
    sql = f"SELECT DISTINCT {columns_str} FROM {table_name} LIMIT {sample_count}"
    rows, _ = spark_db.fetch(sql, conn)
    return [row for row in rows]


def template_command(args):
    """生成模板命令"""
    scenario = args.scenario
    tables = args.tables
    database = getattr(args, 'database', 'default')
    
    print(f"场景: {scenario}, 数据库: {database}")
    print(f"表列表: {tables}")
    
    # 确保数据目录存在
    os.makedirs(data_files, exist_ok=True)
    
    # 创建Spark连接
    conn = spark_db.get_conn('localhost', 0, '', '', database)
    
    try:
        table_names = tables
        wb = Workbook()
        tables_desc = list()
        
        for table_name in table_names:
            print(f"开始处理表 {table_name}...")
            
            # 如果指定了数据库，使用完整表名
            full_table_name = f"{database}.{table_name}" if database != 'default' else table_name
            
            bedrock = aws.get('bedrock-runtime')
            
            schema = get_spark_table_schema(full_table_name, conn)
            schema_search = {item['Name']: item for item in schema}
            
            # 分批处理列（避免超出大模型长度限制）
            batch_count = 40
            begin_index = 0
            pyob_lst = None
            
            while begin_index < len(schema):
                end_index = min(begin_index + batch_count, len(schema))
                columns = schema[begin_index:end_index]
                sample_data = get_spark_sample_data(full_table_name, columns, conn)
                pmt = gen.PROMPT_F.format(full_table_name, columns, sample_data)
                
                questions = [{
                    "role": "user",
                    "content": pmt
                }]
                
                output = llm.query(questions, bedrock)
                print(f"原始输出: {output}")
                
                # 清理输出，移除markdown格式
                cleaned_output = output.strip()
                if cleaned_output.startswith('```json'):
                    cleaned_output = cleaned_output[7:]
                if cleaned_output.endswith('```'):
                    cleaned_output = cleaned_output[:-3]
                cleaned_output = cleaned_output.strip()
                
                print(f"清理后输出: {cleaned_output}")
                
                try:
                    pyob = json.loads(cleaned_output)
                except json.JSONDecodeError as e:
                    print(f"JSON解析错误: {e}")
                    print(f"尝试解析的内容: {cleaned_output}")
                    raise
                
                if pyob_lst:
                    pyob_lst["columns"].extend(pyob["columns"])
                else:
                    pyob_lst = pyob
                    
                begin_index = end_index
            
            # 保存JSON文件（每个表保存独立文件）
            with open(f"{data_files}/{scenario}_{table_name}.json", mode='w') as f:
                output = json.dumps(pyob_lst, ensure_ascii=False, indent=2)
                f.write(output)
            
            # 创建Excel工作表
            ws = wb.create_sheet(title=table_name)
            
            table_desc = pyob_lst['desc']
            columns = pyob_lst['columns']
            tables_desc.append(table_desc)
            
            ws.append(["表名", full_table_name])
            ws.append(["基本信息", table_desc])
            ws.append(["查询规则", ""])
            ws.append(["字段名称", "字段类型", "数据逻辑", "是否保留", "字段含义", "其他说明"])
            
            for column in columns:
                dtype = schema_search[column['name']]['Type']
                dw = f"{column['tips']},{column['option']}"
                desc = column['desc']
                if column['option_desc']:
                    desc += "请注意，值*MUST*只能是这些值中的一个，这些值及它的含义是：" + column['option_desc']
                
                ws.append([column['name'], dtype, dw, "True", desc, ""])
        
        # 创建汇总工作表
        ws = wb.active
        ws.title = "summary"
        ws.append(["场景", scenario])
        ws.append(["场景描述", scenario + "包含这些信息：" + ",".join(tables_desc)])
        ws.append(["查询规则", ""])
        ws.append(["关联规则", ""])
        
        # 保存Excel文件
        wb.save(f"{data_files}/{scenario}.xlsx")
        print(f"处理完毕，生成的template文件位于 {data_files}/{scenario}.xlsx")
        
    finally:
        conn.close()


def prompt_command():
    """生成prompt命令"""
    print("请确认您已经创建了提示词模板并进行了人工Review！默认情况下提示词模板excel文件存放在prompt/data/promptdata/路径下")
    if confirm_action("输入 'y' 表示你创建好了提示词模板并进行了review，输入 'n' 退出："):
        tool.run(data_files, prompt_path)
        aws.upload_file_to_s3(prompt_path, bucket_name, prompt_file)
        aws.upload_file_to_s3(example_path, bucket_name, example_file)
        aws.upload_file_to_s3(rag_path, bucket_name, rag_file)


def main():
    parser = argparse.ArgumentParser(description="Spark数据源提示词生成辅助工具")
    subparsers = parser.add_subparsers(dest="command", help="可用命令")
    
    # template子命令
    parser_template = subparsers.add_parser("template", help="生成提示词模板，供人工review")
    parser_template.add_argument("--scenario", type=str, required=True, help="场景名称")
    parser_template.add_argument("--tables", type=str, nargs="+", required=True, help="数据表列表，表名称之间用空格隔开")
    parser_template.add_argument("--database", type=str, default="default", help="数据库名称，默认为default")
    
    # prompt子命令
    parser_prompt = subparsers.add_parser("prompt", help="生成提示词文件，并上传到S3中")
    
    args = parser.parse_args()
    
    if args.command == "template":
        template_command(args)
    elif args.command == "prompt":
        prompt_command()
    else:
        parser.print_help()


if __name__ == "__main__":
    main()